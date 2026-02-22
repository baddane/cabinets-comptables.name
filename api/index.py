#!/usr/bin/env python3
"""
Backoffice Admin — cabinets-comptables.name
Déploiement Vercel (serverless Python)

Variables à définir dans Vercel Dashboard → Settings → Environment Variables :
  ADMIN_SECRET_KEY     → chaîne aléatoire longue (ex: openssl rand -hex 32)
  ADMIN_USERNAME       → votre login (défaut: admin)
  ADMIN_PASSWORD       → votre mot de passe (obligatoire)
  GITHUB_TOKEN         → Personal Access Token GitHub (scope: repo)
  GITHUB_REPO          → "owner/repo" (ex: "dupont/cabinets-comptables.name")
  GITHUB_BRANCH        → branche cible (défaut: main)
  ANTHROPIC_API_KEY    → clé Anthropic Claude (génération blog)
  GEMINI_API_KEY       → clé Google Gemini (génération blog)
  DEEPSEEK_API_KEY     → clé DeepSeek (génération blog)
  OPENAI_API_KEY       → clé OpenAI ChatGPT (génération blog)
"""

import asyncio
import base64
import csv
import hashlib
import io
import json
import os
import re
import secrets
from datetime import datetime
from pathlib import Path
from typing import Optional

import httpx
from fastapi import Depends, FastAPI, File, Form, HTTPException, Request, UploadFile, status
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from itsdangerous import BadSignature, SignatureExpired, TimestampSigner
from slugify import slugify
from starlette.middleware.base import BaseHTTPMiddleware

# ─── Chemins ──────────────────────────────────────────────────────────────────

# En Vercel, __file__ est dans /var/task/api/index.py → ROOT = /var/task
ROOT      = Path(__file__).parent.parent
ADMIN_DIR = ROOT / "admin"
DATA_FILE = ROOT / "data" / "cabinets.json"
BLOG_FILE = ROOT / "data" / "blog_posts.json"

BLOG_CATEGORIES = ["Conseils", "Fiscalité", "Comptabilité", "Statuts & Juridique"]

LLM_LABELS = {
    "claude":   "Claude (Anthropic)",
    "gemini":   "Gemini (Google)",
    "deepseek": "DeepSeek",
    "chatgpt":  "ChatGPT (OpenAI)",
}

# ─── Config auth ──────────────────────────────────────────────────────────────

TOKEN_TTL_H = 8

# ─── Hachage mot de passe ─────────────────────────────────────────────────────

def _hash_password(password: str, salt: str | None = None) -> str:
    if salt is None:
        salt = secrets.token_hex(16)
    dk = hashlib.pbkdf2_hmac("sha256", password.encode(), salt.encode(), 260_000)
    return f"{salt}${dk.hex()}"


def _verify_password(password: str, stored: str) -> bool:
    if "$" not in stored:
        return False
    salt, _ = stored.split("$", 1)
    return secrets.compare_digest(_hash_password(password, salt), stored)


# ─── Auth chargée depuis les variables d'env Vercel uniquement ───────────────

SECRET_KEY     = os.environ.get("ADMIN_SECRET_KEY") or secrets.token_hex(32)
ADMIN_USERNAME = os.environ.get("ADMIN_USERNAME", "admin")
ANTHROPIC_KEY  = os.environ.get("ANTHROPIC_API_KEY", "")
GEMINI_KEY     = os.environ.get("GEMINI_API_KEY", "")
DEEPSEEK_KEY   = os.environ.get("DEEPSEEK_API_KEY", "")
OPENAI_KEY     = os.environ.get("OPENAI_API_KEY", "")
GOOGLE_API_KEY = os.environ.get("GOOGLE_PLACES_API_KEY", "")

_plain_pw = os.environ.get("ADMIN_PASSWORD", "").strip()
if not _plain_pw:
    # Fallback : hash brut (déconseillé, utiliser ADMIN_PASSWORD)
    ADMIN_PASS_HASH = os.environ.get("ADMIN_PASSWORD_HASH", "")
else:
    ADMIN_PASS_HASH = _hash_password(_plain_pw)

# ─── GitHub API ───────────────────────────────────────────────────────────────

GITHUB_TOKEN  = os.environ.get("GITHUB_TOKEN", "")
GITHUB_REPO   = os.environ.get("GITHUB_REPO", "")   # "owner/repo"
GITHUB_BRANCH = os.environ.get("GITHUB_BRANCH", "main")

_GH_API = "https://api.github.com"
_GH_HEADERS = {
    "Accept": "application/vnd.github+json",
    "X-GitHub-Api-Version": "2022-11-28",
}


async def _github_commit_file(repo_path: str, content_str: str, commit_msg: str) -> bool:
    """
    Commit une mise à jour de fichier via la Git Data API GitHub.
    Supporte les fichiers > 1 MB (contrairement à la Contents API).
    Retourne True si succès.
    """
    if not GITHUB_TOKEN or not GITHUB_REPO:
        return False

    hdrs = {**_GH_HEADERS, "Authorization": f"Bearer {GITHUB_TOKEN}"}

    try:
        async with httpx.AsyncClient(timeout=30) as gh:
            # 1. SHA du dernier commit sur la branche
            r = await gh.get(
                f"{_GH_API}/repos/{GITHUB_REPO}/git/ref/heads/{GITHUB_BRANCH}",
                headers=hdrs,
            )
            r.raise_for_status()
            commit_sha = r.json()["object"]["sha"]

            # 2. Arbre du commit
            r = await gh.get(
                f"{_GH_API}/repos/{GITHUB_REPO}/git/commits/{commit_sha}",
                headers=hdrs,
            )
            r.raise_for_status()
            base_tree_sha = r.json()["tree"]["sha"]

            # 3. Créer un nouveau blob (base64)
            encoded = base64.b64encode(content_str.encode("utf-8")).decode()
            r = await gh.post(
                f"{_GH_API}/repos/{GITHUB_REPO}/git/blobs",
                headers=hdrs,
                json={"content": encoded, "encoding": "base64"},
            )
            r.raise_for_status()
            blob_sha = r.json()["sha"]

            # 4. Créer un nouvel arbre
            r = await gh.post(
                f"{_GH_API}/repos/{GITHUB_REPO}/git/trees",
                headers=hdrs,
                json={
                    "base_tree": base_tree_sha,
                    "tree": [
                        {
                            "path": repo_path,
                            "mode": "100644",
                            "type": "blob",
                            "sha": blob_sha,
                        }
                    ],
                },
            )
            r.raise_for_status()
            new_tree_sha = r.json()["sha"]

            # 5. Créer le commit
            r = await gh.post(
                f"{_GH_API}/repos/{GITHUB_REPO}/git/commits",
                headers=hdrs,
                json={
                    "message": commit_msg,
                    "tree": new_tree_sha,
                    "parents": [commit_sha],
                },
            )
            r.raise_for_status()
            new_commit_sha = r.json()["sha"]

            # 6. Mettre à jour la ref de la branche
            r = await gh.patch(
                f"{_GH_API}/repos/{GITHUB_REPO}/git/refs/heads/{GITHUB_BRANCH}",
                headers=hdrs,
                json={"sha": new_commit_sha},
            )
            r.raise_for_status()
            return True

    except Exception as exc:
        print(f"[github] Erreur commit : {exc}")
        return False


async def _github_get_file_content(repo_path: str) -> str | None:
    """Lit le contenu brut d'un fichier depuis GitHub (toujours à jour, bypass snapshot Vercel)."""
    if not GITHUB_TOKEN or not GITHUB_REPO:
        return None
    hdrs = {
        **_GH_HEADERS,
        "Authorization": f"Bearer {GITHUB_TOKEN}",
        "Accept": "application/vnd.github.raw+json",
    }
    async with httpx.AsyncClient(timeout=30) as gh:
        r = await gh.get(
            f"{_GH_API}/repos/{GITHUB_REPO}/contents/{repo_path}",
            headers=hdrs,
            params={"ref": GITHUB_BRANCH},
        )
        if r.status_code != 200:
            return None
        return r.text


# ─── Google Places API (enrichissement async) ─────────────────────────────────

_PLACES_BASE = "https://maps.googleapis.com/maps/api/place"


def _fmt_phone(raw: str) -> str:
    if not raw:
        return ""
    digits = "".join(c for c in raw if c.isdigit())
    if len(digits) == 11 and digits.startswith("33"):
        digits = "0" + digits[2:]
    if len(digits) == 10 and digits.startswith("0"):
        return " ".join(digits[i:i+2] for i in range(0, 10, 2))
    return raw.strip()


async def _places_find(client: httpx.AsyncClient, cab: dict) -> str | None:
    """Cherche un cabinet sur Google Places, retourne son place_id."""
    params = {
        "query": f"{cab['name']} {cab['city']} expert comptable",
        "type": "accounting",
        "language": "fr",
        "key": GOOGLE_API_KEY,
    }
    if cab.get("lat") and cab.get("lng"):
        params["location"] = f"{cab['lat']},{cab['lng']}"
        params["radius"] = "2000"
    try:
        r = await client.get(f"{_PLACES_BASE}/textsearch/json", params=params)
        results = r.json().get("results", [])
        return results[0].get("place_id") if results else None
    except Exception:
        return None


async def _places_details(client: httpx.AsyncClient, place_id: str) -> dict:
    """Récupère les détails d'un établissement Google Places."""
    params = {
        "place_id": place_id,
        "fields": (
            "formatted_phone_number,international_phone_number,"
            "website,rating,user_ratings_total,"
            "geometry/location,business_status"
        ),
        "language": "fr",
        "key": GOOGLE_API_KEY,
    }
    try:
        r = await client.get(f"{_PLACES_BASE}/details/json", params=params)
        result = r.json().get("result", {})
        return {} if result.get("business_status") == "CLOSED_PERMANENTLY" else result
    except Exception:
        return {}


async def _enrich_one(sem: asyncio.Semaphore, client: httpx.AsyncClient, cab: dict) -> bool:
    """Enrichit un cabinet. Retourne True si au moins un champ a été mis à jour."""
    async with sem:
        place_id = await _places_find(client, cab)
        if not place_id:
            return False
        details = await _places_details(client, place_id)
        if not details:
            return False
        phone = _fmt_phone(
            details.get("formatted_phone_number", "")
            or details.get("international_phone_number", "")
        )
        changed = False
        if phone:
            cab["phone"] = phone
            changed = True
        if details.get("website"):
            cab["website"] = details["website"].strip()
            changed = True
        if details.get("rating") is not None:
            cab["rating"] = str(round(float(details["rating"]), 1))
            changed = True
        if details.get("user_ratings_total"):
            cab["reviews_count"] = details["user_ratings_total"]
            changed = True
        geo = details.get("geometry", {}).get("location", {})
        if geo.get("lat") and geo.get("lng"):
            cab["lat"] = geo["lat"]
            cab["lng"] = geo["lng"]
            changed = True
        return changed


# ─── Données ──────────────────────────────────────────────────────────────────

def load_data() -> list[dict]:
    if DATA_FILE.exists():
        with open(DATA_FILE, encoding="utf-8") as f:
            return json.load(f)
    return []


def load_blog() -> list[dict]:
    if BLOG_FILE.exists():
        with open(BLOG_FILE, encoding="utf-8") as f:
            return json.load(f)
    return []


def data_stats(cabinets: list[dict]) -> dict:
    n = len(cabinets)
    cities  = len({c.get("city", "") for c in cabinets if c.get("city")})
    phones  = sum(1 for c in cabinets if c.get("phone"))
    ratings = sum(1 for c in cabinets if c.get("rating"))
    gps     = sum(1 for c in cabinets if c.get("lat") and c.get("lng"))
    webs    = sum(1 for c in cabinets if c.get("website"))
    return {
        "total": n, "cities": cities,
        "phones": phones,   "pct_phone":  round(phones  / n * 100) if n else 0,
        "ratings": ratings, "pct_rating": round(ratings / n * 100) if n else 0,
        "gps": gps,         "pct_gps":    round(gps     / n * 100) if n else 0,
        "websites": webs,   "pct_web":    round(webs    / n * 100) if n else 0,
    }


# ─── Application ──────────────────────────────────────────────────────────────

app = FastAPI(docs_url=None, redoc_url=None)
templates = Jinja2Templates(directory=str(ADMIN_DIR / "templates"))


class SecurityHeadersMiddleware(BaseHTTPMiddleware):
    async def dispatch(self, request, call_next):
        response = await call_next(request)
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["X-Frame-Options"] = "SAMEORIGIN"
        response.headers["X-XSS-Protection"] = "1; mode=block"
        return response


app.add_middleware(SecurityHeadersMiddleware)

# ─── Utilitaires auth ─────────────────────────────────────────────────────────

_signer = TimestampSigner(SECRET_KEY)


def _make_token(username: str) -> str:
    return _signer.sign(username).decode()


def _get_current_user(request: Request) -> Optional[str]:
    token = request.cookies.get("admin_token")
    if not token:
        return None
    try:
        username = _signer.unsign(token, max_age=TOKEN_TTL_H * 3600)
        return username.decode()
    except (BadSignature, SignatureExpired):
        return None


def _require_auth(request: Request) -> str:
    user = _get_current_user(request)
    if not user:
        raise HTTPException(
            status_code=status.HTTP_303_SEE_OTHER,
            headers={"Location": "/admin/login"},
        )
    return user


def _redirect_login(msg: str = "") -> RedirectResponse:
    url = "/admin/login"
    if msg:
        url += f"?error={msg}"
    return RedirectResponse(url, status_code=303)


# ─── Routes : auth ────────────────────────────────────────────────────────────

@app.get("/admin/login", response_class=HTMLResponse)
async def login_page(request: Request):
    if _get_current_user(request):
        return RedirectResponse("/admin/", 303)
    if not ADMIN_PASS_HASH:
        return HTMLResponse(
            "<h2 style='font-family:sans-serif;color:red;padding:2em'>"
            "⚠️ Variable ADMIN_PASSWORD non définie dans Vercel.<br>"
            "Ajoutez-la dans Vercel → Settings → Environment Variables."
            "</h2>",
            status_code=503,
        )
    return templates.TemplateResponse("login.html", {
        "request": request,
        "error": request.query_params.get("error", ""),
    })


@app.post("/admin/login")
async def login_post(
    request: Request,
    username: str = Form(...),
    password: str = Form(...),
):
    import asyncio
    if username == ADMIN_USERNAME and _verify_password(password, ADMIN_PASS_HASH):
        resp = RedirectResponse("/admin/", 303)
        resp.set_cookie(
            "admin_token", _make_token(username),
            httponly=True, secure=True, samesite="strict",
            max_age=TOKEN_TTL_H * 3600,
        )
        return resp
    await asyncio.sleep(0.5)
    return _redirect_login("Identifiants incorrects")


@app.get("/admin/logout")
async def logout():
    resp = RedirectResponse("/admin/login", 303)
    resp.delete_cookie("admin_token")
    return resp


# ─── Routes : dashboard ───────────────────────────────────────────────────────

@app.get("/admin/", response_class=HTMLResponse)
@app.get("/admin", response_class=HTMLResponse)
async def dashboard(request: Request, user: str = Depends(_require_auth)):
    cabinets = load_data()
    stats    = data_stats(cabinets)
    github_ok = bool(GITHUB_TOKEN and GITHUB_REPO)
    return templates.TemplateResponse("dashboard.html", {
        "request": request, "user": user,
        "stats": stats, "last_gen": None,
        "tasks": {},
        "github_ok": github_ok,
    })


# ─── Routes : cabinets ────────────────────────────────────────────────────────

@app.get("/admin/cabinets", response_class=HTMLResponse)
async def cabinets_list(
    request: Request,
    user: str = Depends(_require_auth),
    q: str = "",
    city: str = "",
    page: int = 1,
):
    cabinets = load_data()
    if q:
        ql = q.lower()
        cabinets = [c for c in cabinets if ql in c.get("name", "").lower()
                                        or ql in c.get("city", "").lower()
                                        or ql in c.get("address", "").lower()]
    if city:
        cabinets = [c for c in cabinets if c.get("city") == city]

    cabinets = sorted(cabinets, key=lambda c: (c.get("city", ""), c.get("name", "")))

    per_page  = 30
    total     = len(cabinets)
    pages     = max(1, (total + per_page - 1) // per_page)
    page      = max(1, min(page, pages))
    sliced    = cabinets[(page - 1) * per_page: page * per_page]

    all_data   = load_data()
    all_cities = sorted({c.get("city", "") for c in all_data if c.get("city")})

    return templates.TemplateResponse("cabinets.html", {
        "request": request, "user": user,
        "cabinets": sliced, "total": total,
        "page": page, "pages": pages,
        "q": q, "city": city,
        "all_cities": all_cities,
    })


@app.get("/admin/cabinets/new", response_class=HTMLResponse)
async def cabinet_new(request: Request, user: str = Depends(_require_auth)):
    return templates.TemplateResponse("cabinet_form.html", {
        "request": request, "user": user,
        "cabinet": {}, "mode": "new", "error": "",
    })


@app.post("/admin/cabinets/new")
async def cabinet_create(
    request: Request, user: str = Depends(_require_auth),
    name: str           = Form(""),
    address: str        = Form(""),
    city: str           = Form(""),
    postal_code: str    = Form(""),
    phone: str          = Form(""),
    website: str        = Form(""),
    rating: str         = Form(""),
    reviews_count: str  = Form("0"),
    rating_info: str    = Form(""),
    category: str       = Form(""),
    open_hours: str     = Form(""),
    lat: str            = Form(""),
    lng: str            = Form(""),
    featured_image: str = Form(""),
    bing_maps_url: str  = Form(""),
    email: str          = Form(""),
    facebook: str       = Form(""),
    instagram: str      = Form(""),
    twitter: str        = Form(""),
    external_id: str    = Form(""),
):
    name = name.strip(); city = city.strip()
    if not name or not city:
        return templates.TemplateResponse("cabinet_form.html", {
            "request": request, "user": user,
            "cabinet": {}, "mode": "new",
            "error": "Le nom et la ville sont obligatoires.",
        })
    new_cab = {
        "name": name, "address": address.strip(),
        "city": city, "postal_code": postal_code.strip(),
        "phone": phone.strip(), "website": website.strip(),
        "rating": rating.strip(), "reviews_count": int(reviews_count or 0),
        "rating_info": rating_info.strip(), "category": category.strip(),
        "open_hours": open_hours.strip(),
        "lat": float(lat) if lat.strip() else "",
        "lng": float(lng) if lng.strip() else "",
        "featured_image": featured_image.strip(), "bing_maps_url": bing_maps_url.strip(),
        "email": email.strip(), "facebook": facebook.strip(),
        "instagram": instagram.strip(), "twitter": twitter.strip(),
        "external_id": external_id.strip(),
    }
    cabinets = load_data()
    cabinets.append(new_cab)
    content = json.dumps(cabinets, ensure_ascii=False, indent=2)
    await _github_commit_file(
        "data/cabinets.json", content,
        f"admin: ajout cabinet {name} ({city})",
    )
    return RedirectResponse("/admin/cabinets?created=1", 303)


@app.get("/admin/cabinets/{idx}/edit", response_class=HTMLResponse)
async def cabinet_edit(request: Request, idx: int, user: str = Depends(_require_auth)):
    cabinets = load_data()
    if idx < 0 or idx >= len(cabinets):
        raise HTTPException(404)
    return templates.TemplateResponse("cabinet_form.html", {
        "request": request, "user": user,
        "cabinet": cabinets[idx], "idx": idx, "mode": "edit", "error": "",
    })


@app.post("/admin/cabinets/{idx}/edit")
async def cabinet_update(
    request: Request, idx: int,
    user: str           = Depends(_require_auth),
    name: str           = Form(""),
    address: str        = Form(""),
    city: str           = Form(""),
    postal_code: str    = Form(""),
    phone: str          = Form(""),
    website: str        = Form(""),
    rating: str         = Form(""),
    reviews_count: str  = Form("0"),
    rating_info: str    = Form(""),
    category: str       = Form(""),
    open_hours: str     = Form(""),
    lat: str            = Form(""),
    lng: str            = Form(""),
    featured_image: str = Form(""),
    bing_maps_url: str  = Form(""),
    email: str          = Form(""),
    facebook: str       = Form(""),
    instagram: str      = Form(""),
    twitter: str        = Form(""),
    external_id: str    = Form(""),
):
    cabinets = load_data()
    if idx < 0 or idx >= len(cabinets):
        raise HTTPException(404)
    name = name.strip(); city = city.strip()
    if not name or not city:
        return templates.TemplateResponse("cabinet_form.html", {
            "request": request, "user": user,
            "cabinet": cabinets[idx], "idx": idx, "mode": "edit",
            "error": "Le nom et la ville sont obligatoires.",
        })
    cabinets[idx].update({
        "name": name, "address": address.strip(),
        "city": city, "postal_code": postal_code.strip(),
        "phone": phone.strip(), "website": website.strip(),
        "rating": rating.strip(), "reviews_count": int(reviews_count or 0),
        "rating_info": rating_info.strip(), "category": category.strip(),
        "open_hours": open_hours.strip(),
        "lat": float(lat) if lat.strip() else "",
        "lng": float(lng) if lng.strip() else "",
        "featured_image": featured_image.strip(), "bing_maps_url": bing_maps_url.strip(),
        "email": email.strip(), "facebook": facebook.strip(),
        "instagram": instagram.strip(), "twitter": twitter.strip(),
        "external_id": external_id.strip(),
    })
    content = json.dumps(cabinets, ensure_ascii=False, indent=2)
    await _github_commit_file(
        "data/cabinets.json", content,
        f"admin: modification cabinet {name} ({city})",
    )
    return RedirectResponse("/admin/cabinets?updated=1", 303)


@app.post("/admin/cabinets/{idx}/delete")
async def cabinet_delete(request: Request, idx: int, user: str = Depends(_require_auth)):
    cabinets = load_data()
    if idx < 0 or idx >= len(cabinets):
        raise HTTPException(404)
    removed = cabinets.pop(idx)
    content = json.dumps(cabinets, ensure_ascii=False, indent=2)
    await _github_commit_file(
        "data/cabinets.json", content,
        f"admin: suppression cabinet {removed.get('name', idx)}",
    )
    return RedirectResponse("/admin/cabinets?deleted=1", 303)


# ─── Routes : blog ────────────────────────────────────────────────────────────

async def _call_llm(llm: str, prompt: str) -> str:
    """Appelle le LLM choisi et retourne le texte brut."""
    if llm == "gemini":
        import google.generativeai as genai
        genai.configure(api_key=GEMINI_KEY)
        model = genai.GenerativeModel("gemini-1.5-flash")
        response = await asyncio.to_thread(model.generate_content, prompt)
        return response.text.strip()
    elif llm == "deepseek":
        from openai import AsyncOpenAI
        client = AsyncOpenAI(api_key=DEEPSEEK_KEY, base_url="https://api.deepseek.com")
        r = await client.chat.completions.create(
            model="deepseek-chat",
            messages=[{"role": "user", "content": prompt}],
            max_tokens=4096,
        )
        return r.choices[0].message.content.strip()
    elif llm == "chatgpt":
        from openai import AsyncOpenAI
        client = AsyncOpenAI(api_key=OPENAI_KEY)
        r = await client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[{"role": "user", "content": prompt}],
            max_tokens=4096,
        )
        return r.choices[0].message.content.strip()
    else:  # claude (défaut)
        from anthropic import AsyncAnthropic
        client = AsyncAnthropic(api_key=ANTHROPIC_KEY)
        message = await client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=4096,
            messages=[{"role": "user", "content": prompt}],
        )
        return message.content[0].text.strip()


@app.get("/admin/blog", response_class=HTMLResponse)
async def blog_list(request: Request, user: str = Depends(_require_auth)):
    posts = load_blog()
    return templates.TemplateResponse("blog.html", {
        "request": request, "user": user,
        "posts": posts,
        "categories": BLOG_CATEGORIES,
        "api_key_ok": bool(ANTHROPIC_KEY or GEMINI_KEY or DEEPSEEK_KEY or OPENAI_KEY),
        "llm_keys": {
            "claude":   bool(ANTHROPIC_KEY),
            "gemini":   bool(GEMINI_KEY),
            "deepseek": bool(DEEPSEEK_KEY),
            "chatgpt":  bool(OPENAI_KEY),
        },
        "llm_labels": LLM_LABELS,
    })


@app.post("/admin/blog/generate")
async def blog_generate(
    request: Request,
    user: str       = Depends(_require_auth),
    topic: str      = Form(""),
    category: str   = Form("Conseils"),
    word_count: str = Form("700"),
    llm: str        = Form("claude"),
):
    llm_key_map = {
        "claude": ANTHROPIC_KEY, "gemini": GEMINI_KEY,
        "deepseek": DEEPSEEK_KEY, "chatgpt": OPENAI_KEY,
    }
    if not llm_key_map.get(llm):
        return JSONResponse(
            {"error": f"Clé API manquante pour {llm}. Définissez la variable dans Vercel → Environment Variables."},
            status_code=400,
        )
    topic = topic.strip()
    if not topic:
        return JSONResponse({"error": "Veuillez saisir un sujet."}, status_code=400)

    # Interlinking : charger les articles existants
    existing_posts = load_blog()
    interlinks = "\n".join(
        f'- <a href="/blog/{p["slug"]}/">{p["title"]}</a>'
        for p in existing_posts
    )
    interlink_section = (
        f"\nArticles existants sur ce site (intégrer 2 à 3 liens pertinents dans le texte) :\n{interlinks}\n"
        "Format du lien interne : <a href=\"/blog/SLUG/\">Titre exact</a>\n"
    ) if interlinks else ""

    prompt = (
        f"Génère un article de blog professionnel en français pour un site annuaire de cabinets comptables.\n\n"
        f"Sujet : {topic}\n"
        f"Catégorie : {category}\n"
        f"Longueur cible : environ {word_count} mots\n"
        f"{interlink_section}\n"
        "Consignes strictes :\n"
        "- Le contenu est en HTML avec UNIQUEMENT ces balises : <p>, <h2>, <ul>, <li>, <strong>, <a>\n"
        "- 3 à 5 sections titrées avec <h2> (pas de <h1>, pas de <h3>)\n"
        "- Termes techniques importants en <strong>\n"
        "- Terminer par un appel à l'action avec ce lien exact : "
        '<a href="/">l\'annuaire des cabinets comptables</a>\n'
        "- Ton professionnel, pratique, orienté entrepreneurs et PME français\n"
        "- Pas de balise <html>, <head>, <body> ni de doctype\n\n"
        "Retourne UNIQUEMENT un objet JSON valide (sans markdown, sans bloc de code) :\n"
        '{"title":"...","slug":"...","description":"...","reading_time":7,"content":"..."}'
    )

    try:
        raw = await _call_llm(llm, prompt)
        if raw.startswith("```"):
            parts = raw.split("```")
            raw = parts[1] if len(parts) > 1 else raw
            if raw.startswith("json"):
                raw = raw[4:]
        raw = raw.strip()
        article = json.loads(raw)
        for field in ("title", "slug", "description", "content"):
            if field not in article:
                raise ValueError(f"Champ manquant : {field}")
        article["category"]     = category
        article["date"]         = datetime.now().strftime("%Y-%m-%d")
        article["reading_time"] = int(article.get("reading_time") or
                                      max(1, len(article["content"].split()) // 200))
        return JSONResponse(article)
    except json.JSONDecodeError as exc:
        return JSONResponse({"error": f"Réponse IA mal formée : {exc}"}, status_code=500)
    except Exception as exc:
        return JSONResponse({"error": str(exc)}, status_code=500)


@app.post("/admin/blog/publish")
async def blog_publish(
    request: Request,
    user: str         = Depends(_require_auth),
    title: str        = Form(""),
    slug: str         = Form(""),
    description: str  = Form(""),
    category: str     = Form("Conseils"),
    reading_time: str = Form("5"),
    date: str         = Form(""),
    content: str      = Form(""),
):
    title = title.strip()
    slug  = slugify(slug.strip()) if slug.strip() else slugify(title)
    if not title or not content:
        return RedirectResponse("/admin/blog?error=missing", 303)

    posts = load_blog()
    existing_slugs = {p["slug"] for p in posts}
    base_slug, counter = slug, 1
    while slug in existing_slugs:
        slug = f"{base_slug}-{counter}"
        counter += 1

    posts.append({
        "slug":         slug,
        "title":        title,
        "description":  description.strip(),
        "date":         date or datetime.now().strftime("%Y-%m-%d"),
        "category":     category,
        "reading_time": max(1, int(reading_time or 5)),
        "content":      content,
    })
    blog_content = json.dumps(posts, ensure_ascii=False, indent=2)
    await _github_commit_file(
        "data/blog_posts.json", blog_content,
        f"admin: publication article « {title} »",
    )
    return RedirectResponse(f"/admin/blog?published={slug}", 303)


@app.post("/admin/blog/{slug}/delete")
async def blog_delete(request: Request, slug: str, user: str = Depends(_require_auth)):
    posts = load_blog()
    removed_title = next((p["title"] for p in posts if p["slug"] == slug), slug)
    posts = [p for p in posts if p["slug"] != slug]
    blog_content = json.dumps(posts, ensure_ascii=False, indent=2)
    await _github_commit_file(
        "data/blog_posts.json", blog_content,
        f"admin: suppression article « {removed_title} »",
    )
    return RedirectResponse("/admin/blog?deleted=1", 303)


# ─── Routes : changement de mot de passe ─────────────────────────────────────

@app.get("/admin/password", response_class=HTMLResponse)
async def password_page(request: Request, user: str = Depends(_require_auth)):
    return templates.TemplateResponse("password.html", {
        "request": request, "user": user, "success": False, "error": "",
        "vercel_mode": True,
    })


@app.post("/admin/password")
async def password_change(
    request: Request, user: str = Depends(_require_auth),
    current: str = Form(""),
    new_pw: str  = Form(""),
    confirm: str = Form(""),
):
    """En mode Vercel, le mot de passe se change dans les env vars Vercel, pas ici."""
    return templates.TemplateResponse("password.html", {
        "request": request, "user": user, "success": False,
        "error": (
            "En mode Vercel, modifiez votre mot de passe directement dans "
            "Vercel → Settings → Environment Variables → ADMIN_PASSWORD, "
            "puis redéployez."
        ),
        "vercel_mode": True,
    })


# ─── Routes : import CSV / Excel ─────────────────────────────────────────────

IMPORT_COLUMNS = [
    ("nom",          "name",          "Nom du cabinet (obligatoire)"),
    ("adresse",      "address",       "Adresse complète (rue, code postal, ville)"),
    ("ville",        "city",          "Ville — extraite automatiquement de l'adresse si absente"),
    ("code_postal",  "postal_code",   "Code postal (ex : 75001)"),
    ("telephone",    "phone",         "Numéro de téléphone"),
    ("site_web",     "website",       "URL du site web"),
    ("note",         "rating",        "Note (ex : 4.5)"),
    ("nb_avis",      "reviews_count", "Nombre d'avis (entier)"),
    ("note_info",    "rating_info",   "Source de la note (ex : Trustpilot (3966))"),
    ("categorie",    "category",      "Catégorie (ex : Comptable)"),
    ("horaires",     "open_hours",    "Horaires d'ouverture"),
    ("latitude",     "lat",           "Latitude GPS (ex : 48.8566)"),
    ("longitude",    "lng",           "Longitude GPS (ex : 2.3522)"),
    ("image",        "featured_image","URL de l'image principale"),
    ("bing_maps",    "bing_maps_url", "URL Bing Maps"),
    ("email",        "email",         "Adresse e-mail de contact"),
    ("facebook",     "facebook",      "URL de la page Facebook"),
    ("instagram",    "instagram",     "URL du profil Instagram"),
    ("twitter",      "twitter",       "URL du profil Twitter / X"),
    ("id_externe",   "external_id",   "Identifiant externe (ex : ypid:...)"),
]

_COL_MAP: dict[str, str] = {}
for _fr, _en, _ in IMPORT_COLUMNS:
    _COL_MAP[_fr.lower()] = _en
    _COL_MAP[_en.lower()] = _en

# Noms anglais supplémentaires (exports Bing Maps / tiers)
_COL_MAP.update({
    "id":            "external_id",
    "emails":        "email",
    "social_medias": "social_medias",
})

_IMPORT_EXAMPLE = [
    "Cabinet Dupont & Associés", "12 rue de la Paix, 75001 Paris", "Paris", "75001",
    "01 23 45 67 89", "https://www.cabinet-dupont.fr", "4.5", "42",
    "Google (42)", "Comptable", "Lun-Ven 09:00-18:00",
    "48.8566", "2.3522", "", "", "contact@cabinet-dupont.fr",
    "", "", "", "",
]


def _extract_city_from_address(address: str) -> tuple[str, str, str]:
    """Extrait (rue, code_postal, ville) depuis une adresse française complète."""
    m = re.search(r",?\s*(\d{4,5})\s+([^,\d]+?)\s*$", address.strip())
    if m:
        street = address[: m.start()].strip().rstrip(",").strip()
        return street, m.group(1).strip(), m.group(2).strip()
    return address, "", ""


def _normalize_import_row(row: dict) -> dict | None:
    out: dict = {}
    for key, val in row.items():
        field = _COL_MAP.get(key.strip().lower().replace(" ", "_"))
        if field:
            out[field] = str(val).strip() if val is not None else ""
    if not out.get("name"):
        return None
    # Auto-extraction ville / code postal depuis adresse complète
    if out.get("address") and not out.get("city"):
        street, postal, city = _extract_city_from_address(out["address"])
        if city:
            out["address"] = street
            if not out.get("postal_code"):
                out["postal_code"] = postal
            out["city"] = city
    if not out.get("city"):
        return None
    # reviews_count depuis rating_info si absent ("Trustpilot (3966)" → 3966)
    if not out.get("reviews_count") and out.get("rating_info"):
        m = re.search(r"\((\d+)\)", out["rating_info"])
        if m:
            out["reviews_count"] = m.group(1)
    try:
        out["reviews_count"] = int(float(out.get("reviews_count") or 0))
    except (ValueError, TypeError):
        out["reviews_count"] = 0
    for f in ("lat", "lng"):
        v = str(out.get(f, "")).replace(",", ".")
        try:
            out[f] = float(v) if v else ""
        except (ValueError, TypeError):
            out[f] = ""
    return out


@app.get("/admin/import", response_class=HTMLResponse)
async def import_page(request: Request, user: str = Depends(_require_auth)):
    return templates.TemplateResponse("import.html", {
        "request": request, "user": user,
        "columns": IMPORT_COLUMNS, "result": None,
        "regen": False,
        "github_saved": None,
        "github_configured": bool(GITHUB_TOKEN and GITHUB_REPO),
        "llm_keys": {
            "claude":   bool(ANTHROPIC_KEY),
            "gemini":   bool(GEMINI_KEY),
            "deepseek": bool(DEEPSEEK_KEY),
            "chatgpt":  bool(OPENAI_KEY),
        },
        "selected_llm": "claude",
    })


@app.get("/admin/import/template.csv")
async def import_template_csv(user: str = Depends(_require_auth)):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c[0] for c in IMPORT_COLUMNS])
    w.writerow(_IMPORT_EXAMPLE)
    return StreamingResponse(
        iter([buf.getvalue().encode("utf-8-sig")]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=modele_cabinets.csv"},
    )


@app.get("/admin/import/template.xlsx")
async def import_template_xlsx(user: str = Depends(_require_auth)):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cabinets"
    for i, (col_name, _, _desc) in enumerate(IMPORT_COLUMNS, 1):
        cell = ws.cell(row=1, column=i, value=col_name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = max(16, len(col_name) + 4)
    for i, val in enumerate(_IMPORT_EXAMPLE, 1):
        ws.cell(row=2, column=i, value=val)
    ws2 = wb.create_sheet("Description colonnes")
    ws2.append(["Colonne", "Description"])
    ws2["A1"].font = Font(bold=True)
    ws2["B1"].font = Font(bold=True)
    for col_name, _, desc in IMPORT_COLUMNS:
        ws2.append([col_name, desc])
    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 50

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=modele_cabinets.xlsx"},
    )


@app.post("/admin/import", response_class=HTMLResponse)
async def import_post(
    request: Request,
    user: str = Depends(_require_auth),
    mode: str = Form("merge"),
    llm: str  = Form("claude"),
    file: UploadFile = File(...),
):
    filename = (file.filename or "").lower()
    result: dict = {"added": 0, "updated": 0, "skipped": 0, "errors": [], "total": 0}
    rows: list[dict] = []

    content = await file.read()

    if filename.endswith(".csv"):
        for enc in ("utf-8-sig", "utf-8", "latin-1"):
            try:
                text = content.decode(enc)
                rows = list(csv.DictReader(io.StringIO(text)))
                break
            except (UnicodeDecodeError, Exception):
                continue
        if not rows:
            result["errors"].append("Impossible de lire le fichier CSV (encodage non reconnu).")
    elif filename.endswith((".xlsx", ".xls")):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            ws = wb.active
            raw_rows = list(ws.iter_rows(values_only=True))
            if raw_rows:
                headers = [str(h or "").strip() for h in raw_rows[0]]
                for row in raw_rows[1:]:
                    rows.append({headers[i]: (str(v) if v is not None else "")
                                 for i, v in enumerate(row)})
        except Exception as e:
            result["errors"].append(f"Erreur lecture Excel : {e}")
    else:
        result["errors"].append("Format non supporté. Utilisez un fichier .csv ou .xlsx")

    github_saved: bool | None = None
    regen = False

    if rows:
        # Lire depuis GitHub (version live) plutôt que le snapshot Vercel figé au déploiement
        if mode == "merge":
            raw_gh = await _github_get_file_content("data/cabinets.json")
            cabinets = json.loads(raw_gh) if raw_gh else load_data()
        else:
            cabinets = []

        idx_extid = {c.get("external_id", ""): i for i, c in enumerate(cabinets) if c.get("external_id")}
        idx_name  = {
            (c.get("name", "").lower(), c.get("city", "").lower()): i
            for i, c in enumerate(cabinets)
        }
        for row_num, raw in enumerate(rows, 2):
            if all(v in ("", "None", None) for v in raw.values()):
                continue
            result["total"] += 1
            cab = _normalize_import_row(raw)
            if cab is None:
                result["errors"].append(f"Ligne {row_num} : nom ou ville manquant — ignorée")
                result["skipped"] += 1
                continue
            if mode == "merge":
                existing_idx = None
                if cab.get("external_id") and cab["external_id"] in idx_extid:
                    existing_idx = idx_extid[cab["external_id"]]
                else:
                    key = (cab["name"].lower(), cab["city"].lower())
                    existing_idx = idx_name.get(key)
                if existing_idx is not None:
                    cabinets[existing_idx].update({k: v for k, v in cab.items() if v != ""})
                    result["updated"] += 1
                else:
                    cabinets.append(cab)
                    new_idx = len(cabinets) - 1
                    if cab.get("external_id"):
                        idx_extid[cab["external_id"]] = new_idx
                    idx_name[(cab["name"].lower(), cab["city"].lower())] = new_idx
                    result["added"] += 1
            else:
                cabinets.append(cab)
                result["added"] += 1

        if result["added"] + result["updated"] > 0:
            data_content = json.dumps(cabinets, ensure_ascii=False, indent=2)
            github_saved = await _github_commit_file(
                "data/cabinets.json", data_content,
                f"admin: import CSV ({result['added']} ajoutés, {result['updated']} mis à jour)",
            )
            if github_saved:
                # Le commit sur data/cabinets.json déclenche automatiquement generate.yml
                regen = True
            else:
                result["errors"].append(
                    "Données non sauvegardées sur GitHub : vérifiez que GITHUB_TOKEN et "
                    "GITHUB_REPO sont configurés dans Vercel → Settings → Environment Variables."
                )
        else:
            github_saved = None  # Rien à sauvegarder

    return templates.TemplateResponse("import.html", {
        "request": request, "user": user,
        "columns": IMPORT_COLUMNS, "result": result,
        "regen": regen,
        "github_saved": github_saved,
        "github_configured": bool(GITHUB_TOKEN and GITHUB_REPO),
        "llm_keys": {
            "claude":   bool(ANTHROPIC_KEY),
            "gemini":   bool(GEMINI_KEY),
            "deepseek": bool(DEEPSEEK_KEY),
            "chatgpt":  bool(OPENAI_KEY),
        },
        "selected_llm": llm.strip().lower(),
    })


@app.post("/admin/import/regenerate")
async def import_regenerate(request: Request, user: str = Depends(_require_auth)):
    """Déclenche manuellement le workflow generate.yml via workflow_dispatch."""
    if not GITHUB_TOKEN or not GITHUB_REPO:
        return JSONResponse(
            {"error": "GITHUB_TOKEN ou GITHUB_REPO non configuré dans Vercel."},
            status_code=400,
        )
    hdrs = {**_GH_HEADERS, "Authorization": f"Bearer {GITHUB_TOKEN}"}
    try:
        async with httpx.AsyncClient(timeout=15) as gh:
            r = await gh.post(
                f"{_GH_API}/repos/{GITHUB_REPO}/actions/workflows/generate.yml/dispatches",
                headers=hdrs,
                json={"ref": GITHUB_BRANCH},
            )
            if r.status_code == 204:
                return JSONResponse({"ok": True})
            return JSONResponse(
                {"error": f"GitHub Actions a répondu {r.status_code} : {r.text[:200]}"},
                status_code=400,
            )
    except Exception as exc:
        return JSONResponse({"error": str(exc)}, status_code=500)


# ─── API JSON ─────────────────────────────────────────────────────────────────

@app.get("/admin/api/stats")
async def api_stats(request: Request, user: str = Depends(_require_auth)):
    return JSONResponse(data_stats(load_data()))


# ─── Outils ───────────────────────────────────────────────────────────────────

@app.get("/admin/outils", response_class=HTMLResponse)
async def outils_page(request: Request, user: str = Depends(_require_auth)):
    return templates.TemplateResponse("scraper.html", {
        "request": request, "user": user,
        "tasks": {},
        "google_key_ok": bool(GOOGLE_API_KEY),
        "vercel_mode": True,
        "github_ok": bool(GITHUB_TOKEN and GITHUB_REPO),
    })


@app.post("/admin/outils/enrich")
async def outils_enrich(
    request: Request,
    user: str  = Depends(_require_auth),
    limit: str = Form("50"),
    cities: str = Form(""),
):
    if not GOOGLE_API_KEY:
        return RedirectResponse("/admin/outils?error=no_key", 303)

    # Lire depuis GitHub pour avoir la version la plus récente (pas le snapshot du déploiement)
    raw = await _github_get_file_content("data/cabinets.json")
    cabinets = json.loads(raw) if raw else load_data()

    city_filter = [c.strip() for c in cities.split(",") if c.strip()] if cities.strip() else None
    safe_limit  = max(1, min(100, int(limit or 50)))

    def needs_enrichment(c: dict) -> bool:
        if c.get("phone") and c.get("rating"):
            return False
        if city_filter and c.get("city") not in city_filter:
            return False
        return True

    all_candidates   = [i for i, c in enumerate(cabinets) if needs_enrichment(c)]
    batch_idx        = all_candidates[:safe_limit]
    total_remaining  = len(all_candidates)

    enriched = 0
    sem = asyncio.Semaphore(5)
    async with httpx.AsyncClient(timeout=15) as client:
        results = await asyncio.gather(
            *[_enrich_one(sem, client, cabinets[i]) for i in batch_idx],
            return_exceptions=True,
        )
        enriched = sum(1 for r in results if r is True)

    if enriched > 0:
        content = json.dumps(cabinets, ensure_ascii=False, indent=2)
        await _github_commit_file(
            "data/cabinets.json", content,
            f"admin: enrichissement Google Places ({enriched} cabinets)",
        )

    remaining_after = max(0, total_remaining - len(batch_idx))
    return RedirectResponse(
        f"/admin/outils?enriched={enriched}&processed={len(batch_idx)}&remaining={remaining_after}",
        303,
    )
